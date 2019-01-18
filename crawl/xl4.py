import openpyxl
from openpyxl.chart import (Reference, BarChart)


book = openpyxl.load_workbook("/Users/lhj/images/xl_output.xlsx")

sheet3 = book.create_sheet()

sheet3.title = '세번째 시트'

rows = [
    ['김일수', 11],
    ['김이수', 22],
    ['김삼수', 33],
    ['김사수', 15],
    ['김오수', 11],
]

for row in rows:
    sheet3.append(row)

datax = Reference(sheet3, min_col=2, 
		min_row=1, max_col=2, max_row=5)

categs = Reference(sheet3, min_col=1,
				 min_row=1, max_row=5)

chart = BarChart()
chart.add_data(data=datax)
chart.set_categories(categs)

chart.legend = None  # 범례
chart.varyColors = True
chart.title = "차트 타이틀"

sheet3.add_chart(chart, "A8")


book.save("/Users/lhj/images/xl_output.xlsx")